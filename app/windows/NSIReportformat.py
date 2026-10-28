"""Formatting helpers for the internal and client NSI Word reports."""

import json
import os
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches


_METHOD_APPENDIX_URL = (
    "https://kktspinetreatment.sharepoint.com/:w:/s/Project-M3/"
    "IQCFw3AhrIqOQbV9j1K3JnU_AXQi--tD8f3wk-ZYdasggBk?e=GZlSTb"
)


def format_value(value, suffix="", digits=3):
    if value is None:
        return "—"
    return f"{value:.{digits}f}{suffix}"


def _add_report_metadata(document, baseline_name, compare_name, offset, client_facing):
    metadata = document.add_paragraph()
    metadata.add_run("Baseline: ").bold = True
    metadata.add_run(baseline_name)
    metadata.add_run("    Comparison: ").bold = True
    metadata.add_run(compare_name)
    if not client_facing:
        metadata.add_run("    Phase offset: ").bold = True
        metadata.add_run(format_value(offset, " s") if offset is not None else "—")


def _add_user_metadata(document, metadata):
    document.add_heading("Report details", level=2)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for cell, text in zip(table.rows[0].cells, ("Field", "Value")):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F2F2")
        cell._tc.get_or_add_tcPr().append(shading)

    for label, value in metadata.items():
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value or "—"


def _add_assessment_summary(
    document, assessment, assessment_reasons, client_facing, client_description
):
    if client_facing:
        document.add_heading("Quality assurance summary", level=2)
        summary_status = document.add_paragraph()
        summary_status.add_run(f"Assessment: {assessment}").bold = True
        document.add_paragraph(client_description)
        document.add_paragraph(
            "Phase 1 and Phase 2 measurements determine the "
            "assessment. Phase 3 results are shown for information "
            "only. See the appendix for how measurements and the "
            "assessment are calculated."
        )
        return

    assessment_line = document.add_paragraph()
    assessment_line.add_run("Device assessment: ").bold = True
    assessment_line.add_run(assessment)
    summary_heading = (
        "Reason for failure"
        if assessment == "FAIL"
        else "Reason for review"
        if assessment == "REVIEW"
        else "Assessment summary"
    )
    document.add_paragraph(summary_heading, style="Heading 2")
    if assessment == "PASS":
        document.add_paragraph(assessment_reasons[0])
    else:
        for reason in assessment_reasons:
            document.add_paragraph(reason, style="List Bullet")
    document.add_paragraph(
        "Assessment uses Phase 1 and Phase 2 only. The selected "
        "baseline is the reference; Phase 3 and correlation "
        "scores are informational. The appendix explains the "
        "calculations and pass, review, and fail rules."
    )


def _add_phase_sections(document, report_sections, client_facing):
    for label, rows, notes, internal_plot, client_plot, boxplot in report_sections:
        phase_plot = client_plot if client_facing else internal_plot
        heading = document.add_heading(label, level=2)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        header_cells = table.rows[0].cells
        for cell, text in zip(header_cells, ("Measurement", "Baseline", "Comparison")):
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F2F2F2")
            cell._tc.get_or_add_tcPr().append(shading)

        for metric, baseline_value, compare_value in rows:
            cells = table.add_row().cells
            for cell, text in zip(cells, (metric, baseline_value, compare_value)):
                cell.text = str(text)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        if phase_plot is not None:
            caption = document.add_paragraph(
                f"{label} waveform comparison (selected time ranges)"
            )
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_picture(BytesIO(phase_plot), width=Inches(6.0))

        if boxplot is not None:
            caption = document.add_paragraph(
                f"{label} measurement distributions. Circles mark "
                "boxplot outliers (1.5×IQR)"
            )
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            document.add_picture(BytesIO(boxplot), width=Inches(6.0))

        if notes:
            notes_paragraph = document.add_paragraph()
            notes_paragraph.add_run(f"Report note ({label}): ").bold = True
            notes_paragraph.add_run(" ".join(notes))


def _add_assessment_criteria(document):
    appendix = document.add_heading(
        "Appendix: Measurement and Assessment Method",
        level=1,
    )
    appendix.alignment = WD_ALIGN_PARAGRAPH.CENTER
    criteria_heading = document.add_heading(
        "Assessment criteria at a glance", level=2
    )
    criteria_heading.runs[0].bold = True
    criteria = document.add_table(rows=1, cols=2)
    criteria.style = "Table Grid"
    criteria.alignment = WD_TABLE_ALIGNMENT.CENTER
    criteria.autofit = True
    for cell, text in zip(criteria.rows[0].cells, ("Result", "Rule")):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F2F2")
        cell._tc.get_or_add_tcPr().append(shading)
    for status, rule in (
        (
            "FAIL",
            "Any Phase 2 cycle-count difference is 20 cycles or "
            "more, or an assessed mean peak-to-peak/period falls "
            "outside the baseline mean ± 3 sample SD.",
        ),
        (
            "REVIEW",
            "No FAIL applies, and at least one Phase 1 cycle count "
            "differs by 2 or more, or a Phase 2 cycle count differs "
            "by at least 20% of its baseline mean but less than "
            "20 cycles.",
        ),
        (
            "INCONCLUSIVE",
            "No FAIL applies, but required comparison data, coverage, "
            "or baseline measurements are missing or insufficient. "
            "This takes precedence over REVIEW.",
        ),
        (
            "PASS",
            "Required data is sufficient; no FAIL or REVIEW "
            "criterion is triggered.",
        ),
    ):
        cells = criteria.add_row().cells
        cells[0].text = status
        cells[1].text = rule
        for run in cells[0].paragraphs[0].runs:
            run.bold = True


def _add_method_appendix_link(document):
    document.add_paragraph(
        "For the detailed formulas, measurement definitions, and "
        "assessment calculation steps, see the "
    )
    link_paragraph = document.paragraphs[-1]
    relationship_id = document.part.relate_to(
        _METHOD_APPENDIX_URL,
        "http://schemas.openxmlformats.org/officeDocument/2006/"
        "relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = "Appendix Measurement and Assessment Method."
    run.append(text)
    hyperlink.append(run)
    link_paragraph._p.append(hyperlink)


def _build_excel_table(metadata, report_sections):
    columns = ["Phase", "Measurement", "Baseline", "Comparison", *metadata]
    values = [
        [label, measurement, baseline, comparison, *metadata.values()]
        for label, rows, *_ in report_sections
        for measurement, baseline, comparison in rows
    ]
    return {
        "columns": columns,
        "graph_payload": {"values": values},
    }


def _add_report_content(
    document,
    baseline_name,
    compare_name,
    offset,
    assessment,
    assessment_reasons,
    client_description,
    report_sections,
    metadata,
    client_facing,
):
    title = document.add_heading(
        "NSI Client Report" if client_facing else "NSI Report",
        level=1,
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_report_metadata(document, baseline_name, compare_name, offset, client_facing)
    _add_user_metadata(document, metadata)
    _add_assessment_summary(
        document,
        assessment,
        assessment_reasons,
        client_facing,
        client_description,
    )
    _add_phase_sections(document, report_sections, client_facing)
    _add_assessment_criteria(document)
    _add_method_appendix_link(document)

    if client_facing:
        closing = document.add_paragraph(
            "Thank you for completing this quality assurance "
            "assessment. We appreciate your time and attention to "
            "the testing process."
        )
        closing.alignment = WD_ALIGN_PARAGRAPH.CENTER


def save_reports(
    path,
    client_path,
    baseline_name,
    compare_name,
    offset,
    assessment,
    assessment_reasons,
    client_description,
    report_sections,
    metadata,
):
    parent = os.path.dirname(os.path.abspath(path))
    if parent:
        os.makedirs(parent, exist_ok=True)

    internal_document = Document()
    _add_report_content(
        internal_document,
        baseline_name,
        compare_name,
        offset,
        assessment,
        assessment_reasons,
        client_description,
        report_sections,
        metadata,
        client_facing=False,
    )
    client_document = Document()
    _add_report_content(
        client_document,
        baseline_name,
        compare_name,
        offset,
        assessment,
        assessment_reasons,
        client_description,
        report_sections,
        metadata,
        client_facing=True,
    )
    internal_document.save(path)
    client_document.save(client_path)


def save_report_json(
    path,
    baseline_name,
    compare_name,
    offset,
    metadata,
    assessment,
    assessment_reasons,
    client_description,
    report_sections,
):
    payload = {
        "report": "NSI Comparison Report",
        "baseline": baseline_name,
        "comparison": compare_name,
        "phase_offset_s": offset,
        "metadata": metadata,
        "assessment": {
            "result": assessment,
            "reasons": assessment_reasons,
            "client_summary": client_description,
        },
        "excel_table": _build_excel_table(metadata, report_sections),
        "sections": [
            {
                "phase": label,
                "columns": ["Measurement", "Baseline", "Comparison"],
                "rows": [
                    {
                        "measurement": measurement,
                        "baseline": baseline,
                        "comparison": comparison,
                    }
                    for measurement, baseline, comparison in rows
                ],
                "notes": notes,
            }
            for (
                label,
                rows,
                notes,
                _internal_plot,
                _client_plot,
                _boxplot,
            ) in report_sections
        ],
    }
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2, allow_nan=False)
        handle.write("\n")
